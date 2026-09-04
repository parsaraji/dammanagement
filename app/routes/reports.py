import io
import openpyxl
from flask import Blueprint, render_template, request, make_response, send_file
from flask_login import login_required
from app.models.reproduction import Calving, Insemination
from app.models.milk import MilkRecord
from app.models.animal import Animal, AnimalStatus
from app.services.jalali import from_jalali, to_jalali
from app.services.pdf_service import generate_pdf_report

bp = Blueprint('reports', __name__)

@bp.route('/reports')
@login_required
def index():
    return render_template('reports/index.html')

def parse_date_safe(raw):
    if not raw or not raw.strip():
        return None
    try:
        return from_jalali(raw.strip())
    except Exception:
        return None

@bp.route('/reports/calving')
@login_required
def calving_report():
    start_d = parse_date_safe(request.args.get('start_date'))
    end_d = parse_date_safe(request.args.get('end_date'))
    fmt = request.args.get('format', 'excel')

    query = Calving.query
    if start_d:
        query = query.filter(Calving.date >= start_d)
    if end_d:
        query = query.filter(Calving.date <= end_d)

    calvings = query.order_by(Calving.date.desc()).all()

    headers = ["تاریخ زایش", "شماره دام مادر", "نوع زایش", "تعداد بره/بزغاله"]
    rows = []
    for c in calvings:
        rows.append([
            to_jalali(c.date),
            c.animal.plastic_tag if c.animal else '---',
            c.calving_type.value if hasattr(c.calving_type, 'value') else str(c.calving_type),
            str(c.offspring_count)
        ])

    if fmt == 'pdf':
        pdf_bytes = generate_pdf_report("گزارش زایش‌ها", headers, rows)
        resp = make_response(pdf_bytes)
        resp.headers['Content-Type'] = 'application/pdf'
        resp.headers['Content-Disposition'] = 'inline; filename=calving_report.pdf'
        return resp
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(r)
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='calving_report.xlsx')

@bp.route('/reports/insemination')
@login_required
def insemination_report():
    fmt = request.args.get('format', 'excel')
    inseminations = Insemination.query.order_by(Insemination.date.desc()).all()

    headers = ["تاریخ تلقیح", "دام ماده", "نوع تلقیح", "اسپرم/پدر", "منجر به آبستنی"]
    rows = []
    for i in inseminations:
        rows.append([
            to_jalali(i.date),
            i.animal.plastic_tag if i.animal else '---',
            i.insemination_type.value if hasattr(i.insemination_type, 'value') else str(i.insemination_type),
            i.sperm.code if i.sperm else (i.sire_animal.plastic_tag if i.sire_animal else '---'),
            'بله' if i.led_to_pregnancy else 'خیر'
        ])

    if fmt == 'pdf':
        pdf_bytes = generate_pdf_report("گزارش تلقیح و آبستنی", headers, rows)
        resp = make_response(pdf_bytes)
        resp.headers['Content-Type'] = 'application/pdf'
        resp.headers['Content-Disposition'] = 'inline; filename=insemination_report.pdf'
        return resp
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(r)
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='insemination_report.xlsx')

@bp.route('/reports/milk')
@login_required
def milk_report():
    fmt = request.args.get('format', 'excel')
    records = MilkRecord.query.order_by(MilkRecord.date.desc()).all()

    headers = ["تاریخ", "شماره دام", "نوبت ۱", "نوبت ۲", "نوبت ۳", "مجموع (kg)"]
    rows = []
    for mr in records:
        rows.append([
            to_jalali(mr.date),
            mr.animal.plastic_tag if mr.animal else '---',
            str(mr.milking1_amount or 0),
            str(mr.milking2_amount or 0),
            str(mr.milking3_amount or 0),
            str(mr.total_amount or 0)
        ])

    if fmt == 'pdf':
        pdf_bytes = generate_pdf_report("گزارش تولید شیر", headers, rows)
        resp = make_response(pdf_bytes)
        resp.headers['Content-Type'] = 'application/pdf'
        resp.headers['Content-Disposition'] = 'inline; filename=milk_report.pdf'
        return resp
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(r)
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='milk_report.xlsx')

@bp.route('/reports/removals')
@login_required
def removals_report():
    fmt = request.args.get('format', 'excel')
    removals = Animal.query.filter(Animal.status.in_([AnimalStatus.READY_FOR_REMOVAL, AnimalStatus.REMOVED])).all()

    headers = ["شماره دام", "تاریخ خروج", "علت خروج", "شماره برگه", "خریدار"]
    rows = []
    for a in removals:
        rows.append([
            a.plastic_tag,
            to_jalali(a.removal_date) if a.removal_date else '---',
            a.removal_reason or '---',
            a.removal_form_number or '---',
            a.removal_buyer_name or '---'
        ])

    if fmt == 'pdf':
        pdf_bytes = generate_pdf_report("گزارش حذف و خروج دام", headers, rows)
        resp = make_response(pdf_bytes)
        resp.headers['Content-Type'] = 'application/pdf'
        resp.headers['Content-Disposition'] = 'inline; filename=removals_report.pdf'
        return resp
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(headers)
        for r in rows:
            ws.append(r)
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name='removals_report.xlsx')
