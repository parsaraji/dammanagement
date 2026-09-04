import io
import openpyxl
from flask import Blueprint, render_template, request, jsonify, redirect, url_for, flash, send_file
from flask_login import login_required, current_user
from app.extensions import db
from app.models.animal import Animal, AnimalStatus, Sex
from app.models.milk import MilkRecord, MilkRecordType
from app.forms.milk import BulkMilkUploadForm
from app.services.jalali import from_jalali, to_jalali
from app.services.milk_service import compute_standardized_milk

bp = Blueprint('milk', __name__)

@bp.route('/milk-bulk')
@login_required
def milk_bulk():
    lactating_females = Animal.query.filter_by(sex=Sex.FEMALE, status=AnimalStatus.ALIVE).order_by(Animal.plastic_tag.asc()).all()
    upload_form = BulkMilkUploadForm()
    return render_template('milk/bulk.html', lactating_females=lactating_females, upload_form=upload_form)

@bp.route('/milk-bulk/submit', methods=['POST'])
@login_required
def milk_bulk_submit():
    raw_date = request.form.get('date')
    d = from_jalali(raw_date) if raw_date else None
    if not d:
        flash('تاریخ رکوردگیری نامعتبر است.', 'danger')
        return redirect(url_for('milk.milk_bulk'))

    rec_type_str = request.form.get('record_type', 'official')
    animal_ids = request.form.getlist('animal_id[]')
    m1_list = request.form.getlist('m1[]')
    m2_list = request.form.getlist('m2[]')
    m3_list = request.form.getlist('m3[]')

    count = 0
    for i, aid in enumerate(animal_ids):
        m1 = float(m1_list[i]) if i < len(m1_list) and m1_list[i] else 0.0
        m2 = float(m2_list[i]) if i < len(m2_list) and m2_list[i] else 0.0
        m3 = float(m3_list[i]) if i < len(m3_list) and m3_list[i] else 0.0
        tot = round(m1 + m2 + m3, 2)

        if tot > 0:
            mr = MilkRecord(
                animal_id=int(aid),
                date=d,
                record_type=MilkRecordType[rec_type_str.upper()],
                milking1_amount=m1,
                milking2_amount=m2,
                milking3_amount=m3,
                total_amount=tot,
                created_by_user_id=current_user.id
            )
            db.session.add(mr)
            count += 1

    db.session.commit()
    flash(f'{count} رکورد شیر با موفقیت ثبت شد.', 'success')
    return redirect(url_for('milk.milk_bulk'))

@bp.route('/milk/download-template')
@login_required
def download_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Milk Records"
    ws.append(["Plastic Tag", "Date (YYYY/MM/DD)", "Milking 1 (kg)", "Milking 2 (kg)", "Milking 3 (kg)"])
    ws.append(["IR-EW-10", to_jalali(request.args.get('date') or None) or "1402/08/15", 1.5, 1.2, 0.0])

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return send_file(
        out,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='milk_record_template.xlsx'
    )

@bp.route('/milk/upload-excel', methods=['POST'])
@login_required
def upload_excel():
    form = BulkMilkUploadForm()
    if form.validate_on_submit():
        f = form.excel_file.data
        wb = openpyxl.load_workbook(f)
        ws = wb.active

        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            plastic_tag = str(row[0]).strip()
            date_str = str(row[1]).strip() if len(row) > 1 and row[1] else None
            m1 = float(row[2]) if len(row) > 2 and row[2] else 0.0
            m2 = float(row[3]) if len(row) > 3 and row[3] else 0.0
            m3 = float(row[4]) if len(row) > 4 and row[4] else 0.0

            animal = Animal.query.filter_by(plastic_tag=plastic_tag).first()
            if animal and date_str:
                d = from_jalali(date_str)
                mr = MilkRecord(
                    animal_id=animal.id,
                    date=d,
                    record_type=MilkRecordType.OFFICIAL,
                    milking1_amount=m1,
                    milking2_amount=m2,
                    milking3_amount=m3,
                    total_amount=round(m1 + m2 + m3, 2),
                    created_by_user_id=current_user.id
                )
                db.session.add(mr)
                count += 1

        db.session.commit()
        flash(f'تعداد {count} رکورد از فایل اکسل با موفقیت پردازش شد.', 'success')
    else:
        flash('خطا در پردازش فایل اکسل.', 'danger')

    return redirect(url_for('milk.milk_bulk'))

@bp.route('/milk/<int:animal_id>/standardize')
@login_required
def standardize(animal_id):
    animal = Animal.query.get_or_404(animal_id)
    res = compute_standardized_milk(animal_id, animal.parity or 1)
    return render_template('milk/standardize.html', animal=animal, res=res)
